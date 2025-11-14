from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import List
from uuid import uuid4

from fastapi import UploadFile


@dataclass
class SlideContent:
    index: int
    text: str


@dataclass
class IngestionResult:
    project_id: str
    stored_path: str
    original_filename: str | None
    slides: List[SlideContent]


class UnsupportedFileFormatError(ValueError):
    """Raised when an unsupported file extension is provided."""


class TextIngestionService:
    """Service that validates, stores and extracts slide text from pitch decks."""

    SUPPORTED_EXTENSIONS = {".pdf", ".pptx"}

    def __init__(self, storage_dir: str | Path = "storage") -> None:
        self.storage_dir = Path(storage_dir)
        self.storage_dir.mkdir(parents=True, exist_ok=True)

    async def ingest(self, project_id: str, upload: UploadFile) -> IngestionResult:
        if not upload.filename:
            raise ValueError("Uploaded file must include a filename.")

        extension = Path(upload.filename).suffix.lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise UnsupportedFileFormatError(
                f"Unsupported file type '{extension}'. Expected one of: {sorted(self.SUPPORTED_EXTENSIONS)}"
            )

        project_dir = self.storage_dir / project_id
        project_dir.mkdir(parents=True, exist_ok=True)

        file_identifier = uuid4().hex
        target_path = project_dir / f"{file_identifier}{extension}"

        file_bytes = await upload.read()
        target_path.write_bytes(file_bytes)

        slides = self._extract_slides(extension=extension, data=file_bytes)

        return IngestionResult(
            project_id=project_id,
            stored_path=str(target_path),
            original_filename=upload.filename,
            slides=slides,
        )

    def _extract_slides(self, extension: str, data: bytes) -> List[SlideContent]:
        if extension == ".pdf":
            return self._extract_pdf(data)
        if extension == ".pptx":
            return self._extract_pptx(data)
        raise UnsupportedFileFormatError(extension)

    def _extract_pdf(self, data: bytes) -> List[SlideContent]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - handled by runtime configuration
            raise RuntimeError("pypdf is required to process PDF files.") from exc

        reader = PdfReader(BytesIO(data))
        slides: List[SlideContent] = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            slides.append(SlideContent(index=index, text=text.strip()))
        return slides

    def _extract_pptx(self, data: bytes) -> List[SlideContent]:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover - handled by runtime configuration
            raise RuntimeError("python-pptx is required to process PPTX files.") from exc

        presentation = Presentation(BytesIO(data))
        slides: List[SlideContent] = []
        for index, slide in enumerate(presentation.slides, start=1):
            collected: List[str] = []
            for shape in slide.shapes:
                text = self._extract_text_from_shape(shape)
                if text:
                    collected.append(text)
            slides.append(SlideContent(index=index, text="\n".join(collected).strip()))
        return slides

    @staticmethod
    def _extract_text_from_shape(shape) -> str:
        text_segments: List[str] = []
        if hasattr(shape, "text") and shape.text:
            text_segments.append(shape.text)
        elif getattr(shape, "has_text_frame", False):
            text_frame = shape.text_frame
            if text_frame and getattr(text_frame, "text", ""):
                text_segments.append(text_frame.text)
            for paragraph in getattr(text_frame, "paragraphs", []):
                if getattr(paragraph, "text", ""):
                    text_segments.append(paragraph.text)
        return "\n".join(segment.strip() for segment in text_segments if segment and segment.strip())

